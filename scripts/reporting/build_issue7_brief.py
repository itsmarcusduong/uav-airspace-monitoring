#!/usr/bin/env python3
"""Ban tom tat #7 de chot voi An (chi tieu, modality, tracker, ke hoach tuan). python scripts/reporting/build_issue7_brief.py --out docs/w3/E1_issue7_chot_chi_tieu_modality_tracker.docx"""
import argparse
from pathlib import Path
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ap = argparse.ArgumentParser(); ap.add_argument("--out", type=Path, required=True); a = ap.parse_args()
d = docx.Document(); s = d.sections[0]; s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(1.6)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman"); st.paragraph_format.space_after = Pt(4)

def p(t, b=False, i=False, size=None, al=None):
    q = d.add_paragraph(); r = q.add_run(t); r.bold = b; r.italic = i
    if size: r.font.size = Pt(size)
    if al is not None: q.alignment = al
    return q
def h(t): p(t, True, size=12.5)
def bl(items):
    for t in items: d.add_paragraph(t, style="List Bullet")
def table(head, rows, w, size=9.5):
    t = d.add_table(rows=1, cols=len(head)); t.style = "Table Grid"
    for i, x in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(x); r.bold = True; r.font.size = Pt(size)
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3"); c._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        cs = t.add_row().cells
        for i, v in enumerate(row):
            cs[i].text = ""; r = cs[i].paragraphs[0].add_run(v); r.font.size = Pt(size)
    for row in t.rows:
        for i, x in enumerate(w): row.cells[i].width = Cm(x)
    d.add_paragraph()

p("E1 — CHỐT ISSUE #7: CHỈ TIÊU, MODALITY, TRACKER VÀ KẾ HOẠCH TUẦN", True, size=14, al=WD_ALIGN_PARAGRAPH.CENTER)
p("Người soạn: Dương Minh Quang · Gửi: Lê Ngọc An · Ngày 21/10/2026 · Issue #7 (W2-06)", i=True, al=WD_ALIGN_PARAGRAPH.CENTER)
p("Mục đích: hai người thống nhất một phương án cho 5 câu hỏi dưới đây, sau đó mang lên GVHD xác nhận và ghi vào ADR-0005. "
  "Mỗi mục có: tình trạng hiện tại, lý do của đề xuất, các phương án và ô để An chọn. Ký hiệu nguồn: [ĐC] Đề cương, [TM] Thuyết minh, [TC] thang chấm cột E1, [KQ] kết quả đã đo trong repo, [SUY LUẬN] nhận định của nhóm.")

h("0. Tình trạng hiện tại (21/10/2026)")
bl(["Dữ liệu: Svanström, 650 video (IR 365, khả kiến 285), 4 lớp; split v1 theo video 454/98/98, test khóa [KQ].",
    "Tuần 3: đã review nhãn 358 video, xuất tập YOLO e1-yolo-v1 (chỉ ảnh khả kiến, 6.869 ảnh train+val); baseline A YOLOv8n@640 đang train trên Kaggle (run 04) [KQ].",
    "Chưa có: chỉ tiêu được GVHD xác nhận, GPU tham chiếu, quyết định modality cuối cùng, tracker chính, thứ tự W5/W6.",
    "Việc chưa chốt kéo theo rủi ro chấm điểm: thang chấm W5 yêu cầu “tích hợp tracker, video test ID ổn định” trong khi Đề cương W5 là augmentation vật thể nhỏ."])

h("1. Chỉ tiêu định lượng")
table(["Nhóm", "Đề xuất hiện tại [TM mục 12]", "Lý do", "Rủi ro"],
      [["Phát hiện", "Drone: P, R ≥ 0,85; mAP@0.5 ≥ 0,80 trên test", "Ngưỡng dùng trong Thuyết minh; đo theo lớp Drone vì đó là mục tiêu chính", "Vật thể nhỏ chiếm đa số (khả kiến: 60,5% hộp Drone < 32×32 px [KQ class_stats.csv]); chưa có số baseline để biết ngưỡng có khả thi không"],
       ["Theo dõi", "HOTA, IDF1, ID switch — chỉ báo khi có GT track hợp lệ", "Svanström không có track ID; mỗi frame ≤ 1 hộp nên có thể dẫn xuất “1 track/video” nhưng đó không phải GT đa đối tượng", "Nếu báo HOTA/IDF1 trên GT dẫn xuất mà không nêu giới hạn → sai về phương pháp"],
       ["Cảnh báo", "Giảm cảnh báo giả ≥ 30% so với chỉ dùng detector; event recall giảm ≤ 5 điểm %", "Đo được trên dữ liệu có sẵn (video âm tính + video Bird/Airplane/Helicopter)", "Cần ≥ 2 giờ video âm tính; review v2 thấy 10 video không có nhãn thực tế vẫn có vật thể bay nên không dùng làm âm tính được [KQ]"],
       ["Thời gian thực", "≥ 15 FPS toàn pipeline; p95 ≤ 500 ms; đo ≥ 3 lần", "Đủ cho hỗ trợ quan sát; video gốc 30 fps", "Chưa có GPU tham chiếu; Kaggle GPU (T4/P100) thay đổi theo phiên"]],
      [2.2, 4.3, 5.0, 5.5], size=9)
p("Gợi ý chọn:", True)
table(["Phương án", "Nội dung", "Ưu", "Nhược"],
      [["1A (đề xuất)", "Giữ ngưỡng như trên, ghi rõ là mục tiêu; sau baseline (W3) điều chỉnh 1 lần qua ADR nếu cần", "Bám Thuyết minh; có mốc rõ để GVHD duyệt", "Có thể không đạt mAP 0,80 với vật thể rất nhỏ"],
       ["1B", "Chưa đặt số; chốt sau khi có baseline", "Không hứa quá", "Không có tiêu chí đạt/không đạt; thang chấm W1 đã yêu cầu chốt metric"],
       ["1C", "Ngưỡng tách theo cỡ vật thể (nhỏ / vừa+lớn)", "Trung thực hơn với vật thể nhỏ", "Thêm việc đo; cần GVHD đồng ý"]],
      [2.4, 6.2, 4.0, 4.4])
p("GPU tham chiếu: đề xuất lấy đúng GPU Kaggle ghi trong env.json của lần đo cuối, công bố tên GPU kèm số FPS. An chọn: ☐ 1A  ☐ 1B  ☐ 1C  ☐ khác: ________")

h("2. Modality (ảnh khả kiến / hồng ngoại)")
p("Tình trạng: baseline A chỉ dùng ảnh khả kiến (ADR-0004). IR có 365 video, độ phân giải 320×256, gần như toàn bộ hộp nhỏ hơn 32×32 px (Bird IR 100%, Drone IR 98,9%) [KQ class_stats.csv].")
table(["Phương án", "Nội dung", "Ưu", "Nhược"],
      [["2A (đề xuất)", "Khả kiến là nhánh chính cho mọi chỉ tiêu; IR là thí nghiệm phụ báo riêng (W7 hoặc W12)", "Bám [TM]; camera sân bay thường là khả kiến; tránh trộn hai phân bố khác nhau trong một con số", "Chỉ dùng 285/650 video cho phần chính"],
       ["2B", "Một model huấn luyện chung khả kiến + IR", "Nhiều dữ liệu hơn", "Khác độ phân giải và kênh màu; khó diễn giải kết quả; tốn GPU"],
       ["2C", "Hai model riêng, báo cả hai", "Đầy đủ", "Gấp đôi thời gian train/đánh giá trong khi lịch đã chặt"]],
      [2.4, 6.2, 4.0, 4.4])
p("An chọn: ☐ 2A  ☐ 2B  ☐ 2C  ☐ khác: ________")

h("3. Tracker")
p("Tình trạng: [ĐC] và [TM] ghi ByteTrack/BoT-SORT; [TC] tuần 5 ghi “ByteTrack/DeepSORT”. Ultralytics hỗ trợ sẵn hai tracker bytetrack.yaml và botsort.yaml (docs.ultralytics.com/modes/track); DeepSORT không có sẵn trong Ultralytics, phải dùng thư viện ngoài.")
bl(["ByteTrack: liên kết cả hộp điểm tin cậy thấp → giữ được vật thể nhỏ, mờ khi điểm tin cậy dao động; nhẹ, không cần mô hình ReID.",
    "BoT-SORT: thêm bù chuyển động camera (GMC) và tùy chọn ReID. Theo bài báo dữ liệu [Svanström 2021], cảm biến đặt trên đế xoay pan/tilt; audit W2 thấy 58 video có bước nhảy tâm hộp lớn, một phần do camera xoay/zoom [KQ] → GMC có thể giúp [SUY LUẬN].",
    "DeepSORT: dùng đặc trưng ngoại hình; với vật thể vài chục pixel, đặc trưng ngoại hình kém tin cậy [SUY LUẬN]; tốn thêm tích hợp ngoài Ultralytics.",
    "Hạn chế chung: không có GT track ID → chỉ đo được ID switch / độ ổn định trên GT dẫn xuất “1 vật thể/video” và nêu rõ giới hạn."])
table(["Phương án", "Nội dung", "Ưu", "Nhược"],
      [["3A (đề xuất)", "ByteTrack là tracker chính; BoT-SORT đối chiếu trên cùng detection", "Bám [ĐC][TM]; cả hai có sẵn, so sánh công bằng", "Lệch chữ “DeepSORT” trong [TC] → cần GVHD đồng ý"],
       ["3B", "ByteTrack chính, DeepSORT đối chiếu", "Bám đúng chữ trong [TC]", "Tích hợp thư viện ngoài; ít phù hợp vật thể nhỏ"],
       ["3C", "Cả ba", "Đủ mọi tài liệu", "Tốn thời gian W5–W6"]],
      [2.4, 6.2, 4.0, 4.4])
p("An chọn: ☐ 3A  ☐ 3B  ☐ 3C  ☐ khác: ________")

h("4. Thứ tự W5/W6 (mâu thuẫn kế hoạch)")
table(["Tuần", "Đề cương", "Thuyết minh", "Thang chấm (E1)"],
      [["W5", "Augmentation vật thể nhỏ", "Tích hợp ByteTrack", "Tích hợp tracker, video test ID ổn định"],
       ["W6", "Tích hợp tracker", "Đối chiếu tracker, demo", "Geofence + cảnh báo theo trajectory"]], [2, 5, 5, 5])
p("Đề xuất 4A: theo thang chấm (thầy chấm theo cột này) — W5 tracker, W6 geofence/cảnh báo; augmentation vật thể nhỏ gộp vào W4 (tuning + augmentation). Phương án 4B: theo Đề cương và giải thích lệch với thầy. An chọn: ☐ 4A  ☐ 4B")

h("5. Việc còn treo liên quan")
bl(["Cặp video IR–khả kiến cùng lớp và số thứ tự: chưa biết có cùng sự kiện không → nếu dùng 2A thì rủi ro này không ảnh hưởng nhánh chính.",
    "GT tracking: đề xuất dẫn xuất “1 track/video” để đo ID switch, ghi giới hạn trong ADR; không báo HOTA/IDF1 như kết quả chính.",
    "Video âm tính cho cảnh báo: cần đếm tổng thời lượng thực có; nếu < 2 giờ thì ghi giới hạn hoặc bổ sung nguồn có license phù hợp."])

h("6. Sau khi An chọn")
bl(["An trả lời trên Issue #7 (comment ghi lựa chọn 1–4).",
    "Quang soạn ADR-0005 theo lựa chọn; An review PR.",
    "Mang bản tóm tắt lên GVHD; ghi kết quả họp (ngày, người dự, kết luận) vào Issue #7 rồi đóng issue."])
a.out.parent.mkdir(parents=True, exist_ok=True); d.save(a.out); print("saved", a.out)
