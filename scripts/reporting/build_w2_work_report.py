#!/usr/bin/env python3
"""Bao cao cong viec tuan W2 (1-2 trang) -> docx. python scripts/reporting/build_w2_work_report.py --out reports/w2/E1_W02_bao_cao_cong_viec.docx"""
import argparse
from pathlib import Path
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ap = argparse.ArgumentParser(); ap.add_argument("--out", type=Path, required=True); a = ap.parse_args()
d = docx.Document()
s = d.sections[0]; s.page_width, s.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"): setattr(s, m, Cm(1.8))
s.top_margin = s.bottom_margin = Cm(1.1)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(2)

def para(text, bold=False, size=None, align=None, italic=False):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    return p

REPO = "https://github.com/itsmarcusduong/uav-airspace-monitoring"
def add_link(par, text, url, size=8.5):
    rid = d.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rp.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rp.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rp.append(sz)
    r.append(rp); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t); h.append(r)
    par._p.append(h)
def plain(par, text, size=8.5):
    r = par.add_run(text); r.font.size = Pt(size)
COMMITS = [("1fd337f", "1fd337f8ae3a23612fe8ce8a984cd55bf387f068"), ("ceb77c4", "ceb77c4042f4df9f995f95eb934d9b41e4109123"),
           ("b609888", "b60988868214aaf9fba20c4279aeb597a3898d56"), ("08647d1", "08647d1674e611c2e64a1f461d2e04165e357191"),
           ("499f101", "499f10192edd3568b2134e9ee665f44d41a3f8ff"), ("20287b2", "20287b27fc9424dcf6b251082445b70651e2cd49"),
           ("2e10c2a", "2e10c2a2973c4b51ba9dc8f693a1e5485d05dfa3"), ("48c8a04", "48c8a047ad1f7383d0a2fbc21c17d5eb231421a2")]
COMMITS2 = [("32682ff", "32682ffeb1fd2181ca2811f5cebf031891d4ee73")]
def table(header, rows, widths, size=8.5):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3"); c._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(size)
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(0)

para("BÁO CÁO CÔNG VIỆC TUẦN 2 — NHÓM E1", True, 13, WD_ALIGN_PARAGRAPH.CENTER)
para("Đề tài: Nghiên cứu và phát triển hệ thống phát hiện, phân loại và theo dõi UAV trong vùng trời sân bay dựa trên học sâu", False, 9.5, WD_ALIGN_PARAGRAPH.CENTER, True)
para("Tuần 2 (04/10–10/10/2026) — Milestone W2. GVHD: TS. Nguyễn Tất Thắng. Sinh viên: Lê Ngọc An, Dương Minh Quang.")
para("1. GitHub: Issue / PR / Commit", True)
t = d.add_table(rows=1, cols=2); t.style = "Table Grid"
for i, h in enumerate(["Hạng mục", "Liên kết"]):
    c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8.5)
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3"); c._tc.get_or_add_tcPr().append(sh)
def row(label):
    cells = t.add_row().cells; cells[0].text = ""; r = cells[0].paragraphs[0].add_run(label); r.font.size = Pt(8.5)
    cells[1].text = ""; return cells[1].paragraphs[0]
p = row("Milestone W2 / Issues W2-00…W2-08")
add_link(p, "Milestones", REPO + "/milestones"); plain(p, " · "); add_link(p, "Issues #1–#9 (W2)", REPO + "/issues"); plain(p, "; W3-01, W3-02 = #10, #11")
p = row("Pull Request W2 evidence")
add_link(p, "PR #13", REPO + "/pull/13"); plain(p, " (dữ liệu, audit, split, ADR); "); add_link(p, "PR #14", REPO + "/pull/14"); plain(p, " (báo cáo, version v0.2) — đều đã merge")
p = row("Commit SHA")
plain(p, "PR #13: ")
for i, (sh7, full) in enumerate(COMMITS):
    add_link(p, sh7, REPO + "/commit/" + full)
    if i < len(COMMITS) - 1: plain(p, ", ")
plain(p, "; PR #14: ")
for sh7, full in COMMITS2: add_link(p, sh7, REPO + "/commit/" + full)
p = row("Review chéo")
plain(p, "Tuần W2 chưa có review chéo trên GitHub: PR #13, #14 do Quang merge. Đã bật bảo vệ nhánh main; từ W3 mọi PR cần An approve trước khi merge.")
p = row("Tag release"); plain(p, "Chưa có (tag đầu tiên tại W4)")
for r_ in t.rows:
    r_.cells[0].width = Cm(6.5); r_.cells[1].width = Cm(10.9)
d.add_paragraph().paragraph_format.space_after = Pt(0)
para("2. Phân chia công việc", True)
table(["Issue", "Nội dung", "Phụ trách", "Reviewer dự kiến", "Trạng thái"],
      [["W2-01/02", "Chạy Kaggle: SHA-256, inventory, EDA", "Quang", "An", "Hoàn thành"],
       ["W2-03", "Tiêu chí audit + xem chéo nhãn", "An (tiêu chí), cả hai (xem)", "Nhau", "Một phần: audit xong; review sơ bộ 358/358, chờ An xác nhận"],
       ["W2-04", "Giao thức split, ADR-0002/0003", "An", "Quang", "Bản nháp (Proposed)"],
       ["W2-05", "Split v1 theo video + rò rỉ", "Quang", "An", "Hoàn thành (chưa đóng băng)"],
       ["W2-07", "Template v0.2 (Mở đầu, 1.1–1.3, 3.2)", "An khung, Quang số liệu", "Nhau", "Bản nháp"],
       ["W2-08", "Báo cáo tuần + PR", "Quang", "An", "Hoàn thành bản nháp"]],
      [1.8, 5.6, 3.4, 2.2, 4.4])
para("3. Planned – Done – Evidence – Problems – Next", True)
table(["Planned", "Done", "Evidence (đường dẫn)", "Problems", "Next week"],
      [["EDA, audit nhãn, split, chuẩn bị baseline",
        "650 video / 203.328 frame, checksum 1.303/1.303 khớp; audit 650 video; split video 454/98/98",
        "results/W02/kaggle_run_01/…/{stage_status.json, class_stats.csv, annotation_audit*.csv, review/review_assignment_w2.csv}; data/splits/video_split_v1*.csv; docs/w2/*.md",
        "253 video lệch số nhãn/frame (231 dư 1 mục), chẩn đoán run 02 không kết luận; 10 video sai nhãn khi xem; 562 cặp video gần giống khác split; chưa họp GVHD",
        "An xác nhận review (W3-01); chốt chính sách căn nhãn (W3-02); export YOLO train/val; train baseline YOLOv8n; duyệt ADR, đóng băng split"]],
      [3.0, 3.6, 4.4, 3.8, 2.6], size=8)
para("4. Kết quả đo (nguồn: CSV/JSON + script)", True)
table(["Chỉ số", "Giá trị", "Nguồn"],
      [["Video / frame", "650 / 203.328 (IR 365, V 285)", "stage_status.json"],
       ["Video theo lớp A/B/D/H", "133 / 130 / 271 / 116", "stage_status.json"],
       ["bbox ngoài ảnh / w,h≤0 / NaN", "0 / 0 / 0", "annotation_audit_summary.json"],
       ["Video số nhãn ≠ số frame", "253 (231 dư đúng 1 mục)", "annotation_audit.csv"],
       ["Split train/val/test (video)", "454 / 98 / 98 (seed 20261217)", "data/splits/video_split_v1_manifest.json"],
       ["Phần cứng chạy", "Kaggle CPU 4 core, không GPU", "env.json"]], [5.2, 6.4, 5.8])
para("5. Dataset card / model card", True)
para("Dataset card: Svanström Drone-detection-dataset (CC0-1.0 theo repo), mục 3.2 của E1_DoAn_v0.2_W02 và data/manifests/. Model card: chưa có (W3).")
para("6. Demo/video", True); para("Chưa có (tuần này không có sản phẩm demo).")
para("7. Vấn đề và rủi ro (trung thực)", True)
for t in ["Split v1 mới ở trạng thái đề xuất; không loại trừ được trùng cảnh/địa điểm giữa các split.",
          "Nhãn lệch 1 mục ở 231 video (7/51 V-Bird có chim nằm ngoài box): tạm loại khỏi huấn luyện cho tới khi xác định offset.",
          "Chưa xác nhận chỉ tiêu và kế hoạch tuần với GVHD (mâu thuẫn Đề cương/Thuyết minh/thang chấm)."]:
    d.add_paragraph(t, style="List Bullet")
a.out.parent.mkdir(parents=True, exist_ok=True); d.save(a.out); print("saved")
