#!/usr/bin/env python3
"""Dien E1_Template_QuyenDoAn_2026_v0.docx -> E1_DoAn_v0.2_W02.docx (giu style template).
    python scripts/reporting/build_w2_project_version.py --template <template.docx> --out reports/w2/E1_DoAn_v0.2_W02.docx
"""
import argparse, copy
from pathlib import Path
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ap = argparse.ArgumentParser()
ap.add_argument("--template", type=Path, required=True)
ap.add_argument("--out", type=Path, required=True)
ap.add_argument("--figdir", type=Path, default=Path("reports/w2/figures"))
ap.add_argument("--kaggle", type=Path, default=Path("results/W02/kaggle_run_01/w2_outputs_bundle"))
a = ap.parse_args()
d = docx.Document(a.template)
P = list(d.paragraphs)
fig_n = [0]; tab_n = [0]

class Cursor:
    def __init__(self, anchor): self.el = anchor._p
    def _put(self, el): self.el.addnext(el); self.el = el
    def para(self, text="", style="Normal", bold=False, italic=False, align=None):
        p = d.add_paragraph(style=style); r = p.add_run(text); r.bold = bold; r.italic = italic
        if align: p.alignment = align
        self._put(p._p); return p
    def bullets(self, items):
        for t in items: self.para(t, style="List Bullet")
    def caption(self, kind, text):
        n = fig_n if kind == "Hình" else tab_n
        n[0] += 1
        p = self.para(f"{kind} {n[0]}. {text}", style="Caption", align=WD_ALIGN_PARAGRAPH.CENTER); return p
    def image(self, path, width_cm=14):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width_cm)); self._put(p._p)
    def table(self, header, rows, widths=None, size=9):
        t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
            sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3"); c._tc.get_or_add_tcPr().append(sh)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""; rr = cells[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths): row.cells[i].width = Cm(w)
        self._put(t._tbl); self.para("", style="Normal"); return t

def fill(idx, builder):
    ph = P[idx]; assert "Nhóm điền" in ph.text, (idx, ph.text)
    ph.text = ""
    cur = Cursor(ph); builder(cur)
    ph._p.getparent().remove(ph._p)

# ---------------- Mở đầu
def mo_dau(c):
    c.para("Bản nháp W2. Nội dung dựa trên Đề cương, Thuyết minh và kết quả Kaggle đã xác minh; các ngưỡng định lượng là đề xuất chờ GVHD xác nhận.", italic=True)
    c.para("Bài toán. Camera mặt đất quan sát vùng trời sân bay cần phát hiện UAV rất nhỏ trong ảnh, phân biệt với chim, máy bay và trực thăng, theo dõi qua nhiều khung hình và báo cho người vận hành khi UAV vào vùng giám sát. Khó khăn chính là vật thể nhỏ (trung vị diện tích bbox hồng ngoại ≈ 180–380 px², khả kiến ≈ 700–4.800 px² theo class_stats.csv) và các lớp dễ nhầm.")
    c.para("Khác bài tập lớn thông thường: đề tài đòi hỏi quy trình dữ liệu có thể tái lập (manifest, checksum, split theo video), so sánh mô hình trên cùng điều kiện, đánh giá phát hiện–theo dõi–cảnh báo bằng số liệu, và đóng gói mã nguồn/minh chứng theo từng tuần trên GitHub.")
    c.image(a.figdir / "fig_overview.png", 15.5)
    c.caption("Hình", "Sơ đồ tổng quan bài toán và chu trình xử lý đề xuất.")
    c.para("Mục tiêu định lượng (đề xuất theo Thuyết minh mục 12, chờ GVHD chốt ở ADR-0002):", bold=True)
    c.table(["Mục tiêu", "Chỉ số", "Ngưỡng đề xuất", "Tập đo"],
        [["Phát hiện UAV", "Precision, Recall, mAP@0.5 lớp Drone", "P, R ≥ 0,85; mAP@0.5 ≥ 0,80", "Test (khóa tới W11)"],
         ["Theo dõi", "HOTA/IDF1/ID switch", "Chỉ báo khi có GT track hợp lệ (Svanström không có track ID)", "Chưa xác định"],
         ["Cảnh báo", "Giảm cảnh báo giả; event recall", "Giảm ≥ 30% so với detector-only; lệch recall ≤ 5 điểm %", "Video âm tính ≥ 2 giờ"],
         ["Thời gian thực", "FPS, p95 độ trễ", "≥ 15 FPS; p95 ≤ 500 ms (GPU tham chiếu = GPU Kaggle, đo ≥ 3 lần)", "Cùng phần cứng"]],
        [3.2, 4.2, 5.6, 3.2])
    c.caption("Bảng", "Mục tiêu định lượng đề xuất (chưa được GVHD xác nhận).")
    c.para("Phân chia công việc (từ W2; chi tiết theo Issue ở báo cáo tuần):", bold=True)
    c.table(["Thành viên", "Vai trò", "Phần phụ trách W2", "Review chéo"],
        [["Dương Minh Quang", "Lead B: dữ liệu, tích hợp, kiểm thử", "Chạy Kaggle (W2-01/02), split v1 (W2-05), báo cáo W2 (W2-08), nền tảng GitHub (W2-00)", "An review PR"],
         ["Lê Ngọc An", "Lead A: phân tích, thiết kế, thí nghiệm", "Tiêu chí audit (W2-03), giao thức split/ADR (W2-04), khung Mở đầu–Chương I (W2-07)", "Quang review PR"]],
        [3.2, 4.0, 6.8, 2.2])
    c.caption("Bảng", "Phân chia công việc W2. Chỉ ghi đóng góp đã có Issue/commit; số Issue/PR điền ở báo cáo tuần.")
    c.para("Cấu trúc báo cáo: Chương I giới thiệu; II cơ sở lý thuyết; III yêu cầu, dữ liệu và phương pháp; IV thiết kế hệ thống; V cài đặt và thực nghiệm; VI kết luận.")

# ---------------- 1.1
def s11(c):
    c.para("UAV nhỏ xuất hiện trong ảnh với rất ít điểm ảnh, hình dạng gần giống chim, và chịu ảnh hưởng của nền trời, mây, ánh sáng ngược, nén video. Vì vậy mô hình tốt trên ảnh thông thường chưa đủ để kết luận cho giám sát sân bay; cần dữ liệu quan sát từ mặt đất và đánh giá theo video.")
    c.para("Các dữ liệu công khai khác nhau về góc nhìn. Bộ Svanström và cộng sự dùng camera cố định mặt đất, có UAV, chim, máy bay, trực thăng ở dạng video khả kiến và hồng ngoại [11]. VisDrone [4, 5] và UAVDT [7] có camera đặt trên UAV nhìn xuống nên chỉ phù hợp làm tham khảo về vật thể nhỏ và benchmark, không phải mục tiêu UAV bay trên trời. Về theo dõi, ByteTrack [2] và BoT-SORT [3] là hai phương pháp liên kết đa đối tượng được dùng làm cơ sở; RT-DETR [6] là hướng phát hiện thời gian thực để đối chiếu, còn Ultralytics [1] cung cấp giao diện detector–tracker thống nhất.")
    c.table(["Dữ liệu", "Góc nhìn", "Mục tiêu UAV bay", "Vai trò trong đề tài"],
        [["Svanström [11]", "Camera cố định mặt đất (kiểm chứng: 650 video, 203.328 frame)", "Có (Drone 271 video) + 3 lớp gây nhầm", "Dữ liệu chính (CC0-1.0 theo repo; CHƯA đóng băng)"],
         ["VisDrone [4, 5]", "Camera trên UAV nhìn xuống", "Không phải mục tiêu chính", "Chỉ tham khảo"],
         ["UAVDT [7]", "Camera trên UAV", "Không phải mục tiêu chính", "Chỉ tham khảo"],
         ["Anti-UAV", "Chưa kiểm tra đầy đủ", "NEEDS_REVIEW", "Ứng viên bổ sung; license media chưa xác minh, chưa tải"]],
        [3.0, 5.0, 3.8, 4.4])
    c.caption("Bảng", "So sánh các nguồn dữ liệu. Nguồn: tài liệu W1 và kiểm chứng Kaggle W2; hàng Anti-UAV chưa xác minh.")
    c.para("Giới hạn của dữ liệu chính (cần nói rõ): chỉ có ban ngày, thu ở số ít vị trí tại Thụy Điển, không đại diện sân bay Việt Nam, mưa, sương hay ban đêm. Vì vậy kết quả của đồ án là kết quả trên benchmark công khai, không phải khẳng định vận hành thực tế.")

def s12(c):
    c.para("Mục tiêu tổng quát: xây nguyên mẫu xử lý video camera mặt đất để phát hiện, phân loại và theo dõi UAV, kèm cảnh báo hỗ trợ người quan sát và bộ đánh giá tái lập được. Câu hỏi kỹ thuật: (Q1) cấu hình detector nào cân bằng độ chính xác/tốc độ cho vật thể nhỏ; (Q2) thông tin theo thời gian (tracking) giảm cảnh báo giả bao nhiêu; (Q3) hiệu năng thay đổi thế nào theo cự ly và modality.")
    c.table(["Mục tiêu", "Chỉ số", "Ngưỡng đề xuất", "Tuần", "Evidence"],
        [["Dữ liệu sạch, tái lập", "Checksum 1.303 file; split theo video", "mismatch = 0; không rò rỉ trực tiếp", "W2", "checksum_verification.csv; video_split_v1.csv"],
         ["Detector baseline", "mAP@0.5, P/R theo lớp, theo cỡ bbox", "Theo mục tiêu trên", "W3–W4", "Chưa có (W3)"],
         ["Tracker", "ID switch, độ ổn định ID", "Chưa có GT track hợp lệ", "W5", "Chưa có"],
         ["Cảnh báo", "False alarm / event recall", "−30% / ≤ 5 điểm %", "W6–W7", "Chưa có"],
         ["Tốc độ", "FPS, p95 latency", "≥ 15 FPS; ≤ 500 ms", "W9", "Chưa có"]],
        [3.2, 3.8, 3.6, 1.6, 4.2], size=8.5)
    c.caption("Bảng", "Mục tiêu → chỉ số → ngưỡng đề xuất → minh chứng. Các dòng \"Chưa có\" chưa được thực hiện.")

def s13(c):
    c.para("Đối tượng: UAV/Drone và ba lớp gây nhầm (Bird, Airplane, Helicopter) trong video camera mặt đất. Phạm vi MVP: đọc video tệp, phát hiện + phân loại 4 lớp, theo dõi, polygon ROI trên ảnh và cảnh báo có xác nhận theo thời gian; đánh giá offline.")
    c.table(["MVP (bắt buộc)", "Mở rộng (nếu còn thời gian)"],
        [["Detector baseline + 1 cấu hình đối chứng; ByteTrack; ROI + cảnh báo; benchmark cố định", "BoT-SORT đối chiếu, RT-DETR, luồng camera thử nghiệm, Docker, PostgreSQL"]],
        [8.2, 8.0])
    c.caption("Bảng", "Phạm vi MVP và mở rộng.")
    c.para("Giới hạn: không đại diện điều kiện đêm/mưa/sương và sân bay Việt Nam; polygon là vùng trên ảnh, không phải vùng cấm bay pháp lý; hệ thống hỗ trợ quan sát, không xác định ý định của đối tượng; mọi số liệu nêu trong báo cáo phải truy được tới file kết quả.")

def s14(c):
    c.para("(Bản nháp, hạn 17/10.) Kết quả dự kiến: quy trình dữ liệu tái lập (manifest, checksum, split theo video); hai cấu hình detector so sánh trên cùng điều kiện; tích hợp tracker và cảnh báo theo ROI; phân tích lỗi với hard negatives; mã nguồn, model card, dataset card và benchmark có thể chạy lại. Đóng góp chính là quy trình đánh giá theo video kèm các giới hạn dữ liệu được nêu rõ (căn nhãn, rò rỉ cảnh, thiếu track ID).")

def s15(c):
    c.para("(Bản nháp, hạn 17/10.) Chương I giới thiệu; Chương II nền tảng và công trình liên quan; Chương III yêu cầu, dữ liệu, phương pháp, thiết kế thí nghiệm; Chương IV thiết kế hệ thống; Chương V cài đặt và thực nghiệm; Chương VI kết luận và hướng phát triển.")

# ---------------- 3.2
def s32(c):
    K = a.kaggle
    c.para("Nguồn: Drone-detection-dataset của Svanström và cộng sự [11], bản tải thủ công 04/10/2026; license CC0-1.0 theo repository (kiểm tra tại W1). Toàn vẹn dữ liệu trên Kaggle: 1.303/1.303 file khớp SHA-256 với bản W1 (checksum_verification.csv, stage_status.json).")
    c.table(["Thông tin", "Giá trị (nguồn)"],
        [["Số video có nhãn", "650: IR 365 (320×256), Visible 285 (640×512), 30 fps (video_inventory.csv)"],
         ["Số frame giải mã", "203.328 (stage_status.json)"],
         ["Lớp (video)", "Airplane 133, Bird 130, Drone 271, Helicopter 116"],
         ["Nhãn", "MATLAB MCOS *_LABELS.mat, đọc bằng mcos-decoder; mỗi frame ≤ 1 bbox (x, y, w, h); không có track ID"],
         ["Metadata", "Video_dataset_description.xlsx: cảm biến, lớp, số thứ tự, khoảng cách (CLOSE 133 / MEDIUM 292 / DISTANT 225), loại UAV/nguồn"],
         ["Phần cứng chạy EDA", "Kaggle CPU 4 core, Python 3.13, OpenCV 4.14; không GPU (env.json)"]],
        [3.8, 12.4])
    c.caption("Bảng", "Dataset card tóm tắt.")
    c.image(K / "fig_videos_by_modality_class.png", 12)
    c.caption("Hình", "Số video theo modality và lớp (Kaggle run 01).")
    c.para("Kích thước bbox (diện tích, px²; nhỏ = dưới 32×32 = 1.024 px²) theo class_stats.csv:", bold=True)
    c.table(["Modality", "Lớp", "Video", "Frame", "Tỷ lệ bbox nhỏ", "Diện tích trung vị"],
        [["IR", "Airplane", 74, "23.243", "96,4%", "197"], ["IR", "Bird", 79, "24.854", "100%", "178"],
         ["IR", "Drone", 157, "49.043", "98,9%", "219"], ["IR", "Helicopter", 55, "16.955", "86,3%", "381"],
         ["V", "Airplane", 59, "18.585", "45,2%", "1.269"], ["V", "Bird", 51, "16.118", "82,3%", "702"],
         ["V", "Drone", 114, "35.496", "60,5%", "900"], ["V", "Helicopter", 61, "19.034", "22,7%", "4.804"]],
        [2.4, 3.0, 2.2, 2.6, 3.2, 3.2])
    c.caption("Bảng", "Thống kê kích thước bbox. Vật thể nhỏ chiếm ưu thế, đặc biệt ở hồng ngoại.")
    c.image(K / "fig_bbox_area_by_class.png", 14)
    c.caption("Hình", "Phân bố diện tích bbox theo lớp và modality (Kaggle run 01).")
    c.para("Chất lượng nhãn (audit tự động toàn bộ 650 video; chi tiết docs/w2/annotation_audit_w2.md):", bold=True)
    c.bullets(["Không có bbox ngoài ảnh, w/h ≤ 0 hoặc NaN.",
               "12 video không có bbox dương (giữ làm âm tính).",
               "253 video có số mục nhãn ≠ số frame; 231 video dư đúng 1 mục (hệ thống theo nhóm V-Drone, V-Bird, IR-Bird). Chưa biết nhãn lệch đầu hay cuối chuỗi; chẩn đoán bằng độ tương phản (Kaggle run 02) không kết luận được → tạm loại nhóm này khỏi huấn luyện/đánh giá chính (ADR-0003, đề xuất).",
               "58 video có nhảy tâm bbox lớn; ít nhất một video (IR_AIRPLANE_018) do camera dịch chuyển/zoom, nên giả định \"camera cố định\" không đúng cho toàn bộ dữ liệu.",
               "Review bằng mắt sơ bộ 358/358 video (3 frame/video, có hỗ trợ công cụ): 204 OK, 144 không rõ, 10 sai nhãn (7 ở V-Bird); chưa có xác nhận của người thứ hai."])
    c.para("Giao thức chia train/val/test (đề xuất v1, chưa đóng băng):", bold=True)
    c.para("Đơn vị chia là video (mp4 và nhãn luôn cùng split), phân tầng modality × lớp, tỷ lệ 70/15/15, seed chọn bằng tìm kiếm tất định (cơ sở 20261011, chọn 20261217) để cân bằng thêm theo frame, bbox và khoảng cách. Không chia theo frame. Tăng cường chỉ trên train; ngưỡng chọn trên val; test khóa tới đánh giá cuối.")
    c.table(["Split", "Video", "Frame", "Frame có bbox"],
        [["train", 454, "142.187", "136.013"], ["val", 98, "30.545", "28.943"], ["test", 98, "30.596", "29.596"]], [3, 3, 4, 4])
    c.caption("Bảng", "Split v1 (data/splits/video_split_v1.csv).")
    c.image(a.figdir / "fig_split_v1_by_class.png", 11)
    c.caption("Hình", "Số video theo lớp trong từng split v1.")
    c.para("Nguy cơ rò rỉ còn lại: (i) 562 cặp video khác split có khung hình gần giống theo dHash (437 cặp cùng lớp) — dữ liệu thu ở số ít vị trí nên cảnh lặp; (ii) 123/279 cặp IR/V cùng lớp+số thứ tự nằm khác split, nhưng chưa có bằng chứng chúng là cùng sự kiện nên chưa gom; (iii) workbook không có thông tin phiên quay. Do đó không tuyên bố test là \"cảnh chưa thấy\". Cả hai là sàng lọc thô, không phải kết luận.")

fill(57, mo_dau); fill(63, s11); fill(66, s12); fill(69, s13); fill(72, s14); fill(75, s15); fill(96, s32)

# abbreviations table
t = d.tables[2]
extra = [("mAP", "mean Average Precision"), ("IR", "Infrared - hồng ngoại"), ("ROI", "Region of Interest - vùng giám sát"),
         ("HOTA", "Higher Order Tracking Accuracy"), ("IDF1", "ID F1 score"), ("UAV", "Unmanned Aerial Vehicle - máy bay không người lái"),
         ("ADR", "Architecture Decision Record - ghi nhận quyết định")]
last = t.rows[-1]; last.cells[0].text = extra[0][0]; last.cells[1].text = extra[0][1]
for k, v in extra[1:]:
    new = copy.deepcopy(last._tr); t._tbl.append(new)
    row = t.rows[-1]; row.cells[0].text = k; row.cells[1].text = v

# references
for i, p in enumerate(d.paragraphs):
    if p.text.startswith("[10] NVIDIA"):
        cur = Cursor(p)
        r11 = cur.para("[11] Svanström, F., Alonso-Fernandez, F., Englund, C. Drone-detection-dataset. https://github.com/DroneDetectionThesis/Drone-detection-dataset (truy cập 04/10/2026).")
        break
a.out.parent.mkdir(parents=True, exist_ok=True)
d.save(a.out); print("saved", a.out)
