#!/usr/bin/env python3
"""Bao cao cong viec tuan W3 (1-2 trang) -> docx. python scripts/reporting/build_w3_work_report.py --out reports/w3/E1_W03_bao_cao_cong_viec.docx"""
import argparse
from pathlib import Path
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ap = argparse.ArgumentParser(); ap.add_argument("--out", type=Path, required=True); a = ap.parse_args()
d = docx.Document()
s = d.sections[0]; s.page_width, s.page_height = Cm(21), Cm(29.7)
for m in ("left_margin", "right_margin"): setattr(s, m, Cm(1.8))
s.top_margin = s.bottom_margin = Cm(1.1)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(2)

def para(text, bold=False, size=None, align=None, italic=False):
    p = d.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    return p

REPO = "https://github.com/itsmarcusduong/uav-airspace-monitoring"
def add_link(par, text, url, size=8.5):
    rid = d.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rp = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rp.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rp.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rp.append(sz)
    r.append(rp); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t); h.append(r)
    par._p.append(h)
def plain(par, text, size=8.5):
    r = par.add_run(text); r.font.size = Pt(size)
def table(header, rows, widths, size=8.5):
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3"); c._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(v); r.font.size = Pt(size)
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(0)


PRS = [(22, "0c42206", "0c42206249534dc07c73ffc2dd81b9829ee72af2", "job xuất YOLO, train, ADR-0004 (#16)", [("d686117", "d6861170e18773e9afaf85a3f7725719f21fe2de"), ("bf1cfe1", "bf1cfe181bfbef3a9201f2a10d07e950c364449c")]),
       (23, "680043b", "680043bf9cfd43033a92688cdaabbd547049c6cc", "review nhãn v2, notebook Kaggle (#10)", [("40dcdb2", "40dcdb2b24676b7bc817663dd5ec794404d31767"), ("d204db3", "d204db393b29cb33d19c9a864dde10d583b64cbf")]),
       (25, "e762257", "e762257", "bằng chứng review v2, e1-yolo-v1 (#11, #16)", [("e143e66", "e143e66ec61522a02510f202fe60544d7491d1c9"), ("b619a2e", "b619a2ec6e476e9cf9c3e102e2f072ae23e05f61")]),
       (26, "220cd8a", "220cd8a", "job phân tích val, khung model card (#19)", [("98f69c8", "98f69c8e681bf06e5360838c17ddc64eddb3448b")]),
       (28, "41a82fe", "41a82fe21f4aba52a42d7bcf10ef7ca1435f8198", "sửa theo review của An (#19)", [("39e8ff0", "39e8ff0a4f43a01a9c1c16e7b04d37683e8041aa"), ("4f1e3dd", "4f1e3ddcbab96695803c32e34624b52a02c8ad36")]),
       (29, "a54cbec", "a54cbec81689d4ecbac78e9e41932867276f36fb", "ADR-0005 (#7)", [("276c2c2", "276c2c27639eae721c2791342126f148ec2c8e25")]),
       (31, "4532858", "45328584402d6f13011d1812df138c142908481e", "An: license Ultralytics (#20)", [("99d3534", "99d353479a9923f4a33984114c192d44b517347a")]),
       (32, "47eec28", "47eec2877c33bf04fe11ed384a017e5f36b72f47", "An: giao thức train (#17)", [("fbea796", "fbea7963bc9b5416dd2cc7beeb278f1b10aa138b"), ("d98603f", "d98603faa52f0b82db55ab69c295a6a9aafd86e0")]),
       (33, "9597fce", "9597fce22219073880f0ff8a47ba16c7cbd39a21", "đóng gói dữ liệu zip (#18)", [("9c90ba2", "9c90ba20241dbe9bbc279e1c8760eca8fa761a75")]),
       (34, "0b4c79a", "0b4c79afa2a2abb27c304c1eeb7494f01344c3fe", "bằng chứng run 04/05, model card (#18, #19)", [("fc71d2a", "fc71d2afd696a6af68f5f454286375e0288b4513"), ("2f678b8", "2f678b86112fa1ba171350b1ef225b240bb166dc"), ("e073188", "e073188801fa59ab165c56947cb89ffc89219bfe")])]
PENDING = "[chờ run 04]"

para("BÁO CÁO CÔNG VIỆC TUẦN 3 — NHÓM E1", True, 13, WD_ALIGN_PARAGRAPH.CENTER)
para("Đề tài: Nghiên cứu và phát triển hệ thống phát hiện, phân loại và theo dõi UAV trong vùng trời sân bay dựa trên học sâu", False, 9.5, WD_ALIGN_PARAGRAPH.CENTER, True)
para("Tuần 3 (11/10–17/10/2026) — Milestone W3: Train baseline (config, weights, mAP/PR). GVHD: TS. Nguyễn Tất Thắng. Sinh viên: Lê Ngọc An, Dương Minh Quang.")
para("1. GitHub: Issue / PR / Commit", True)
t = d.add_table(rows=1, cols=2); t.style = "Table Grid"
for i, h in enumerate(["Hạng mục", "Liên kết"]):
    c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8.5)
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3"); c._tc.get_or_add_tcPr().append(sh)
def row(label):
    cells = t.add_row().cells; cells[0].text = ""; r = cells[0].paragraphs[0].add_run(label); r.font.size = Pt(8.5)
    cells[1].text = ""; return cells[1].paragraphs[0]
p = row("Milestone W3 / Issues")
add_link(p, "Milestones", REPO + "/milestones"); plain(p, " · "); add_link(p, "Issues #10, #11, #16–#21", REPO + "/issues?q=milestone%3A%22W3+%E2%80%93+Train+baseline+%28config%2C+weights%2C+mAP%2FPR%29%22")
for num, m7, mfull, desc, commits in PRS:
    p = row(f"PR #{num}")
    add_link(p, f"PR #{num}", f"{REPO}/pull/{num}"); plain(p, f" — {desc}. Commit: ")
    for i, (sh7, full) in enumerate(commits):
        add_link(p, sh7, f"{REPO}/commit/{full}")
        if i < len(commits) - 1: plain(p, ", ")
    plain(p, "; merge "); add_link(p, m7, f"{REPO}/commit/{mfull}")
p = row("Review chéo")
plain(p, "An review #22, #23, #25, #26 (sau merge, 3 lỗi → sửa ở #28), approve #28 sau 1 lần yêu cầu sửa, merge #28, #29, #31; Quang review và approve #31, #32 của An. #34 merge trước khi An review lại (cần kịp báo cáo), An review sau.")
p = row("Kaggle (COMMIT dùng cho run)"); add_link(p, "680043b", REPO + "/commit/680043bf9cfd43033a92688cdaabbd547049c6cc"); plain(p, " run 03a/03b; "); add_link(p, "9597fce", REPO + "/commit/9597fce22219073880f0ff8a47ba16c7cbd39a21"); plain(p, " run 03b bản 2, 04; "); add_link(p, "4532858", REPO + "/commit/45328584402d6f13011d1812df138c142908481e"); plain(p, " run 05")
p = row("Tag release"); plain(p, "Chưa có (tag đầu tiên tại W4)")
for r_ in t.rows:
    r_.cells[0].width = Cm(4.6); r_.cells[1].width = Cm(12.8)
d.add_paragraph().paragraph_format.space_after = Pt(0)
para("2. Phân chia công việc", True)
table(["Issue", "Nội dung", "Phụ trách", "Review", "Trạng thái"],
      [["#10, #11", "Review nhãn v2 358 video; chính sách căn nhãn (độ nhạy IoU + thăm dò)", "Quang", "An", "Đóng (PR #23, #25)"],
       ["#16", "ADR-0004 + xuất e1-yolo-v1 (6.869 ảnh)", "Quang", "An", "Đóng (PR #22, #25, #33)"],
       ["#17", "Giao thức train + config baseline A", "An", "Quang", "Đóng (PR #32)"],
       ["#18", "Train YOLOv8n@640, Kaggle T4; weights Drive + SHA-256", "Quang", "An", "Đóng (PR #34)"],
       ["#19", "Phân tích val (run 05) + model card", "Quang, An", "Nhau", "Đóng (PR #26, #28, #34)"],
       ["#20", "License Ultralytics/YOLOv8n", "An", "Quang", "Đóng (PR #31)"],
       ["#21", "Báo cáo tuần + version v0.3", "Quang, An", "Nhau", "Bản nộp (PR #27 + PR báo cáo cuối)"],
       ["#7", "Chỉ tiêu, modality, tracker", "Cả hai", "—", "Nhóm chốt (ADR-0005, PR #29); chờ GVHD"]],
      [1.6, 6.6, 1.9, 1.3, 6.0])
para("3. Planned – Done – Evidence – Problems – Next", True)
table(["Planned", "Done", "Evidence (đường dẫn trong repo)", "Problems", "Next week"],
      [["Chốt nhãn, xuất dữ liệu YOLO, train baseline A, đánh giá val, model card",
        "Review v2 358 video; e1-yolo-v1 216 video / 6.869 ảnh; baseline A train xong (36 epoch, best 21); phân tích val, chọn ngưỡng 0,25; model card",
        "results/W03/review/; kaggle_run_03b*/; kaggle_run_04/train_evidence/; kaggle_run_05/analysis/; docs/w3/model_card_baseline_A.md; ADR-0004, ADR-0005",
        "Số liệu trên val (lạc quan); mAP@0.5:0.95 Drone 0,266; val không có ảnh nền trống thật; chiều lệch nhãn chưa xác định; lần đầu gắn input Kaggle lỗi",
        "Kiểm chứng offset (CVAT #24 + so IoU), cấu hình đối chứng B + augmentation, tag release W4"]],
      [3.0, 3.8, 4.4, 3.8, 2.4], size=8)
para("4. Kết quả đo (nguồn: CSV/JSON + script)", True)
table(["Chỉ số (validation)", "Giá trị", "Nguồn"],
      [["mAP@0.5 / mAP@0.5:0.95 (4 lớp)", "0,807 / 0,342", "kaggle_run_04/train_evidence/baseline_A_val_metrics.json"],
       ["Drone: AP@0.5 / AP@0.5:0.95", "0,809 / 0,266", "baseline_A_val_per_class.csv"],
       ["Ngưỡng vận hành (F1 Drone lớn nhất)", "conf 0,25: Drone P 0,832, R 0,882, F1 0,856", "kaggle_run_05/analysis/operating_points.csv"],
       ["Recall hộp nhỏ / vừa / lớn (conf 0,25)", "0,863 / 0,803 / 0,903", "kaggle_run_05/analysis/recall_by_size.csv"],
       ["Train", "36/50 epoch, best epoch 21, ≈ 27 phút, 1 × T4, ultralytics 8.4.183", "train/results.csv, env.json"],
       ["Dữ liệu train / val", "5.667 / 1.202 ảnh, fingerprint d0ac7e23…", "kaggle_run_03b*/export_evidence/dataset_card.json"]], [5.4, 5.6, 6.4])
para("5. Dataset card / model card", True)
para("Dataset card: results/W03/kaggle_run_03b/export_evidence/dataset_card.json. Model card: docs/w3/model_card_baseline_A.md (đủ run 04/05). Weights: Google Drive weights/W03 (best.pt, last.pt), SHA-256 và link trong weights_manifest.json.")
para("6. Demo/video", True); para("Chưa có (tuần này không có sản phẩm demo).")
para("7. Vấn đề và rủi ro (trung thực)", True)
for t_ in ["Toàn bộ số đo trên val, val vừa dùng để dừng sớm và chọn ngưỡng → lạc quan; chưa so với chỉ tiêu trên test (ADR-0005).",
           "Mục nhãn rỗng của video dư 1 mục nằm ở frame 0 trong khi vật thể đã xuất hiện → nghiêng về offset1; nếu đúng, dữ liệu phải xuất lại và train lại ở W4.",
           "Val không có ảnh nền trời trống: chưa đo được cảnh báo giả (W7). Một seed, một lần chạy: chưa có khoảng dao động.",
           "Chỉ tiêu, GPU tham chiếu và thứ tự W5/W6 (ADR-0005) chưa được GVHD xác nhận."]:
    d.add_paragraph(t_, style="List Bullet")
a.out.parent.mkdir(parents=True, exist_ok=True); d.save(a.out); print("saved")
