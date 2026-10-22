#!/usr/bin/env python3
"""Dien E1_Template_QuyenDoAn_2026_v0.docx -> E1_DoAn_v0.3_W03.docx (giu style template).

Them so voi v0.2: muc luc / danh sach hinh / danh sach bang bang truong TOC that, chu thich hinh-bang
bang truong SEQ (chu thich bang dat TREN bang, hinh DUOI hinh), muc 1.4-1.5 day du, muc 3.2 cap nhat
ket qua W3, tai lieu tham khao theo IEEE. Mo bang Word: References -> Update Table (hoac F9).
    python scripts/reporting/build_w3_project_version.py --template <template.docx> --out reports/w3/E1_DoAn_v0.3_W03.docx \
        --fig-w2 reports/w2/figures --fig-w3 reports/w3/figures --kaggle-w2 results/W02/kaggle_run_01/w2_outputs_bundle
"""
import argparse
import copy
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ap = argparse.ArgumentParser()
ap.add_argument("--template", type=Path, required=True)
ap.add_argument("--out", type=Path, required=True)
ap.add_argument("--fig-w2", type=Path, required=True)
ap.add_argument("--fig-w3", type=Path, required=True)
ap.add_argument("--kaggle-w2", type=Path, required=True)
a = ap.parse_args()
d = docx.Document(a.template)
P = list(d.paragraphs)


def field_run(p, instr, placeholder=""):
    """Chen truong Word (begin/instr/separate/text/end) vao doan p."""
    def r(child):
        run = OxmlElement("w:r"); run.append(child); p._p.append(run); return run
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r(b)
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = f" {instr} "; r(t)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r(s)
    x = OxmlElement("w:t"); x.set(qn("xml:space"), "preserve"); x.text = placeholder; r(x)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r(e)


class Cursor:
    n = {"Hình": 0, "Bảng": 0}

    def __init__(self, anchor):
        self.el = anchor._p

    def _put(self, el):
        self.el.addnext(el); self.el = el

    def para(self, text="", style="Normal", bold=False, italic=False, align=None):
        p = d.add_paragraph(style=style)
        if text:
            r = p.add_run(text); r.bold = bold; r.italic = italic
        if align is not None:
            p.alignment = align
        self._put(p._p); return p

    def rich(self, parts, style="Normal"):
        """parts = [(text, bold)]"""
        p = d.add_paragraph(style=style)
        for t, b in parts:
            r = p.add_run(t); r.bold = b
        self._put(p._p); return p

    def bullets(self, items):
        for t in items:
            self.para(t, style="List Bullet")

    def heading3(self, text):
        return self.para(text, style="Heading 3")

    def caption(self, kind, text):
        Cursor.n[kind] += 1
        p = d.add_paragraph(style="Caption"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"{kind} ")
        field_run(p, f"SEQ {kind} \\* ARABIC", str(Cursor.n[kind]))
        p.add_run(f". {text}")
        self._put(p._p); return p

    def image(self, path, width_cm, caption):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm)); self._put(p._p)
        self.caption("Hình", caption)

    def table(self, caption, header, rows, widths=None, size=9, note=None):
        cp = self.caption("Bảng", caption); cp.paragraph_format.keep_with_next = True
        t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
            sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "D9E2F3"); c._tc.get_or_add_tcPr().append(sh)
        trp = t.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement("w:tblHeader"); trp.append(hdr)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""; rr = cells[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self._put(t._tbl)
        if note:
            q = self.para(note, italic=True); q.runs[0].font.size = Pt(8.5)
        else:
            self.para("")
        return t


def fill(idx, builder):
    ph = P[idx]; assert "Nhóm điền" in ph.text, (idx, ph.text)
    ph.text = ""; cur = Cursor(ph); builder(cur); ph._p.getparent().remove(ph._p)


def replace_with_field(idx, instr, hint):
    p = P[idx]; assert p.text.startswith("["), (idx, p.text)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    field_run(p, instr, hint)


# ============================================================ MỞ ĐẦU
def mo_dau(c):
    c.para("Phiên bản v0.3 (tuần 3). Số liệu dẫn từ các tệp kết quả trong kho mã nguồn (đường dẫn ghi kèm); các ngưỡng định lượng là đề xuất, chờ giảng viên hướng dẫn xác nhận.", italic=True)
    c.para("Bài toán. Camera đặt trên mặt đất quan sát vùng trời quanh sân bay cần phát hiện UAV chỉ chiếm vài chục đến vài trăm điểm ảnh, phân biệt chúng với chim, máy bay và trực thăng, theo dõi qua nhiều khung hình và báo cho người vận hành khi UAV đi vào vùng giám sát. Khó khăn chính là vật thể nhỏ (trung vị diện tích hộp bao ở ảnh hồng ngoại khoảng 180–380 px², ở ảnh khả kiến khoảng 700–4.800 px², theo class_stats.csv) và các lớp có hình dạng dễ nhầm.")
    c.para("Điểm khác với một bài tập lớn thông thường là đề tài phải có quy trình dữ liệu tái lập được (manifest, checksum, chia dữ liệu theo video), so sánh mô hình trong cùng điều kiện, đánh giá phát hiện – theo dõi – cảnh báo bằng số liệu và lưu minh chứng theo từng tuần trên GitHub.")
    c.image(a.fig_w2 / "fig_overview.png", 15.5, "Sơ đồ tổng quan bài toán và chu trình xử lý đề xuất.")
    c.para("Mục tiêu định lượng đề xuất theo Thuyết minh mục 12 được tóm tắt trong Bảng 1.")
    c.table("Mục tiêu định lượng đề xuất (chưa được giảng viên hướng dẫn xác nhận).",
            ["Mục tiêu", "Chỉ số", "Ngưỡng đề xuất", "Tập đo"],
            [["Phát hiện UAV", "Precision, Recall, mAP@0.5 lớp Drone", "P, R ≥ 0,85; mAP@0.5 ≥ 0,80", "Test (khóa tới đánh giá cuối)"],
             ["Theo dõi", "HOTA, IDF1, ID switch", "Chỉ báo khi có nhãn track hợp lệ (Svanström không có track ID)", "Chưa xác định"],
             ["Cảnh báo", "Tỷ lệ cảnh báo giả; event recall", "Giảm ≥ 30% so với chỉ dùng detector; recall giảm ≤ 5 điểm %", "Video âm tính ≥ 2 giờ"],
             ["Thời gian thực", "FPS, độ trễ p95", "≥ 15 FPS; p95 ≤ 500 ms trên GPU tham chiếu, đo ≥ 3 lần", "Cùng phần cứng"]],
            [3.2, 4.2, 5.6, 3.2])
    c.para("Phân chia công việc được ghi theo Issue trên GitHub; Bảng 2 liệt kê phần đã có Issue hoặc commit tới hết tuần 3.")
    c.table("Phân chia công việc tới tuần 3 (chỉ liệt kê đóng góp đã có Issue/commit).",
            ["Thành viên", "Vai trò", "Đóng góp chính (Issue)", "Review chéo"],
            [["Dương Minh Quang", "Lead B: dữ liệu, tích hợp, kiểm thử",
              "W2: Kaggle EDA/audit (#2, #3), split v1 (#6), báo cáo (#9). W3: review nhãn v2 (#10), kiểm tra lệch nhãn (#11), xuất dữ liệu YOLO (#16), huấn luyện baseline A (#18), báo cáo (#21)", "An"],
             ["Lê Ngọc An", "Lead A: phân tích, thiết kế, thí nghiệm",
              "W2: tiêu chí audit (#4), giao thức split và ADR (#5), khung Mở đầu – Chương I (#8). W3: kiểm tra license Ultralytics/YOLOv8n (#20), giao thức huấn luyện (#17), phân tích kết quả và model card (#19)", "Quang"]],
            [3.0, 3.6, 7.6, 2.0])
    c.para("Cấu trúc báo cáo gồm sáu chương, được tóm tắt ở mục 1.5.")


# ============================================================ 1.1 – 1.3 (giữ nội dung v0.2, chỉnh câu chữ)
def s11(c):
    c.para("UAV cỡ nhỏ xuất hiện trong ảnh với rất ít điểm ảnh, có hình dạng gần giống chim và chịu ảnh hưởng của nền trời, mây, ánh sáng ngược và nén video. Vì vậy một mô hình cho kết quả tốt trên ảnh thông thường chưa đủ để kết luận cho bài toán giám sát vùng trời sân bay; cần dữ liệu quan sát từ mặt đất và cách đánh giá theo video.")
    c.para("Các bộ dữ liệu công khai khác nhau về góc nhìn. Bộ dữ liệu của Svanström và cộng sự dùng camera đặt trên mặt đất, gồm UAV, chim, máy bay và trực thăng ở dạng video khả kiến và hồng ngoại [11]. VisDrone [4], [5] và UAVDT [7] dùng camera gắn trên UAV nhìn xuống nên chỉ phù hợp làm tham khảo về vật thể nhỏ và cách xây dựng benchmark. Về theo dõi, ByteTrack [2] và BoT-SORT [3] là hai phương pháp liên kết đa đối tượng làm cơ sở; RT-DETR [6] là hướng phát hiện thời gian thực để đối chiếu; Ultralytics [1], [12] cung cấp giao diện thống nhất giữa detector và tracker.")
    c.table("So sánh các nguồn dữ liệu (nguồn: tài liệu W1 và kiểm chứng Kaggle W2).",
            ["Dữ liệu", "Góc nhìn", "Mục tiêu UAV bay", "Vai trò trong đề tài"],
            [["Svanström [11]", "Camera mặt đất (kiểm chứng: 650 video, 203.328 frame)", "Có (Drone 271 video) + 3 lớp gây nhầm", "Dữ liệu chính (CC0-1.0 theo repository)"],
             ["VisDrone [4], [5]", "Camera trên UAV nhìn xuống", "Không phải mục tiêu chính", "Chỉ tham khảo"],
             ["UAVDT [7]", "Camera trên UAV", "Không phải mục tiêu chính", "Chỉ tham khảo"],
             ["Anti-UAV", "Chưa kiểm tra đầy đủ", "Cần xem xét", "Ứng viên bổ sung; license dữ liệu chưa xác minh, chưa tải"]],
            [3.0, 5.0, 3.8, 4.4])
    c.para("Giới hạn của dữ liệu chính: chỉ có điều kiện ban ngày, thu tại số ít vị trí ở Thụy Điển, không đại diện cho sân bay Việt Nam, mưa, sương hay ban đêm. Do đó kết quả của đồ án là kết quả trên benchmark công khai, không phải khẳng định về vận hành thực tế.")


def s12(c):
    c.para("Mục tiêu tổng quát: xây dựng nguyên mẫu xử lý video từ camera mặt đất để phát hiện, phân loại và theo dõi UAV, kèm cảnh báo hỗ trợ người quan sát và bộ đánh giá tái lập được. Ba câu hỏi kỹ thuật: (Q1) cấu hình detector nào cân bằng tốt nhất giữa độ chính xác và tốc độ với vật thể nhỏ; (Q2) thông tin theo thời gian (tracking) giảm cảnh báo giả được bao nhiêu; (Q3) hiệu năng thay đổi thế nào theo cự ly và modality.")
    c.table("Mục tiêu → chỉ số → ngưỡng đề xuất → minh chứng (cập nhật tới tuần 3).",
            ["Mục tiêu", "Chỉ số", "Ngưỡng đề xuất", "Tuần", "Minh chứng"],
            [["Dữ liệu sạch, tái lập", "Checksum 1.303 tệp; split theo video", "Không sai khác; không rò rỉ trực tiếp", "W2–W3", "checksum_verification.csv; video_split_v1.csv; dataset_card.json"],
             ["Detector baseline", "mAP@0.5, P/R theo lớp, theo cỡ bbox", "Như Bảng 1", "W3–W4", "W3: mAP@0.5 = 0,807 trên val (Kaggle run 04, mục 5.3); test chưa đánh giá"],
             ["Tracker", "ID switch, độ ổn định ID", "Chưa có nhãn track hợp lệ", "W5", "Chưa có"],
             ["Cảnh báo", "Cảnh báo giả / event recall", "−30% / ≤ 5 điểm %", "W6–W7", "Chưa có"],
             ["Tốc độ", "FPS, độ trễ p95", "≥ 15 FPS; ≤ 500 ms", "W9", "Chưa có"]],
            [3.2, 3.8, 3.6, 1.6, 4.2], size=8.5)


def s13(c):
    c.para("Đối tượng: UAV và ba lớp gây nhầm (chim, máy bay, trực thăng) trong video từ camera mặt đất. Phạm vi MVP: đọc video tệp, phát hiện và phân loại 4 lớp, theo dõi, vùng giám sát dạng đa giác trên ảnh và cảnh báo có xác nhận theo thời gian; đánh giá offline.")
    c.table("Phạm vi MVP và phần mở rộng.", ["MVP (bắt buộc)", "Mở rộng (nếu còn thời gian)"],
            [["Detector baseline + 1 cấu hình đối chứng; ByteTrack; ROI + cảnh báo; benchmark cố định",
              "BoT-SORT đối chiếu, RT-DETR, tối ưu bằng TensorRT [10], luồng camera thử nghiệm, Docker, PostgreSQL"]],
            [8.2, 8.0])
    c.para("Giới hạn: không đại diện điều kiện đêm, mưa, sương và sân bay Việt Nam; đa giác ROI là vùng trên ảnh, không phải vùng cấm bay pháp lý; hệ thống chỉ hỗ trợ quan sát, không suy đoán ý định của đối tượng; mọi số liệu trong báo cáo phải truy được tới tệp kết quả.")


# ============================================================ 1.4
def s14(c):
    c.para("Kết quả dự kiến của đồ án gồm bốn nhóm sản phẩm: dữ liệu, mô hình, phần mềm và tài liệu nghiên cứu. Mỗi sản phẩm có định danh phiên bản (commit, SHA-256 hoặc tag) để người khác có thể kiểm tra và chạy lại. Bảng 6 liệt kê các sản phẩm và trạng thái tại ngày 22/10/2026.")
    c.table("Danh sách sản phẩm (artifact) có phiên bản và trạng thái tại tuần 3.",
            ["Sản phẩm", "Loại", "Định danh / vị trí", "Trạng thái"],
            [["Manifest và checksum dữ liệu gốc", "Dữ liệu", "data/manifests/; checksum_verification.csv (1.303 tệp)", "Hoàn thành (W1–W2)"],
             ["Split theo video v1 (454/98/98)", "Dữ liệu", "data/splits/video_split_v1.csv, SHA-256 7e3994…", "Đề xuất, chưa đóng băng"],
             ["Review nhãn v2 (358 video)", "Dữ liệu", "results/W03/review/review_assignment_w3.csv", "Hoàn thành (W3)"],
             ["Tập huấn luyện YOLO e1-yolo-v1", "Dữ liệu", "dataset_card.json, fingerprint d0ac7e23…", "Hoàn thành (W3)"],
             ["Detector baseline A (YOLOv8n @640)", "Mô hình", "best.pt SHA-256 65dd67fb… (Drive) + model card (docs/w3/model_card_baseline_A.md)", "Hoàn thành (W3), đánh giá trên val"],
             ["Detector đối chứng B", "Mô hình", "Tag release W4", "Kế hoạch (W4)"],
             ["Mô-đun theo dõi (ByteTrack, đối chiếu BoT-SORT)", "Phần mềm", "W5–W6", "Kế hoạch"],
             ["Mô-đun ROI và cảnh báo theo quỹ đạo", "Phần mềm", "W6–W7", "Kế hoạch"],
             ["Pipeline đầu-cuối, API, dashboard, nhật ký", "Phần mềm", "Tag release W8", "Kế hoạch"],
             ["Benchmark FPS/độ trễ cố định", "Nghiên cứu", "Tag release W11", "Kế hoạch"],
             ["Báo cáo, model card, dataset card, hướng dẫn tái lập", "Tài liệu", "Tag release W14", "Đang cập nhật hằng tuần"]],
            [5.0, 2.0, 6.2, 3.0], size=8.5)
    c.para("Đóng góp của đồ án được xác định ở mức kỹ thuật, có thể kiểm chứng bằng số liệu:")
    c.bullets([
        "(C1) Quy trình chuẩn bị và đánh giá dữ liệu theo video cho bài toán phát hiện UAV từ mặt đất: chia dữ liệu theo video để tránh rò rỉ, sàng lọc cảnh gần trùng, đo ảnh hưởng của lệch nhãn một khung hình bằng IoU giữa các khung liên tiếp, và review nhãn có quy tắc ghi rõ.",
        "(C2) So sánh có kiểm soát hai cấu hình detector (YOLOv8n @640 và một cấu hình đối chứng) trên cùng split, cùng ngân sách huấn luyện, có phân tích theo lớp và theo cỡ vật thể, dùng PyTorch [9] và Ultralytics [12].",
        "(C3) Đánh giá tác động của thông tin thời gian (tracking và xác nhận nhiều khung hình) lên tỷ lệ cảnh báo giả, trong giới hạn dữ liệu không có nhãn track ID.",
        "(C4) Nguyên mẫu phần mềm đầu-cuối và bộ minh chứng tái lập (mã nguồn, cấu hình, SHA-256, tag release)."])
    c.para("Đóng góp của từng thành viên được ghi theo Issue và commit trên GitHub, tổng hợp ở Bảng 2 và cập nhật trong báo cáo công việc hằng tuần. Hai thành viên review chéo Pull Request của nhau; báo cáo ghi trung thực các tuần chưa có review.")


# ============================================================ 1.5
def s15(c):
    c.para("Báo cáo gồm sáu chương. Bảng 7 nêu câu hỏi mà mỗi chương trả lời; các đoạn dưới đây tóm tắt nội dung.")
    c.table("Câu hỏi mà mỗi chương trả lời.", ["Chương", "Câu hỏi trả lời"],
            [["I. Giới thiệu", "Bài toán là gì, vì sao cần giải quyết, mục tiêu và phạm vi đến đâu?"],
             ["II. Cơ sở lý thuyết", "Những phương pháp phát hiện, theo dõi và đánh giá nào làm nền tảng cho giải pháp?"],
             ["III. Yêu cầu, dữ liệu, phương pháp", "Hệ thống cần đáp ứng gì, dữ liệu được chuẩn bị và kiểm soát ra sao, thí nghiệm được thiết kế thế nào?"],
             ["IV. Thiết kế hệ thống", "Các mô-đun được tổ chức và kết nối như thế nào?"],
             ["V. Cài đặt và thực nghiệm", "Kết quả đo được là bao nhiêu, lỗi nằm ở đâu?"],
             ["VI. Kết luận", "Đồ án đạt được gì, còn hạn chế gì, nên phát triển tiếp theo hướng nào?"]],
            [5.0, 11.2])
    for t in [
        "Chương I trình bày bối cảnh giám sát UAV quanh sân bay, các nguồn dữ liệu công khai và giới hạn của chúng, mục tiêu định lượng, phạm vi MVP, sản phẩm dự kiến và đóng góp.",
        "Chương II tóm tắt kiến thức nền về phát hiện vật thể một giai đoạn (YOLO) và dựa trên Transformer (RT-DETR), theo dõi đa đối tượng theo hướng tracking-by-detection (ByteTrack, BoT-SORT), các độ đo mAP, HOTA, IDF1, và so sánh các công trình liên quan.",
        "Chương III phân tích yêu cầu chức năng và phi chức năng, mô tả dữ liệu Svanström cùng quy trình kiểm định, chia dữ liệu theo video và xử lý lệch nhãn, phương pháp đề xuất và giao thức thí nghiệm.",
        "Chương IV trình bày kiến trúc hệ thống từ đọc video, phát hiện, theo dõi, cảnh báo theo ROI tới API, dashboard và nhật ký; thiết kế dữ liệu và các rủi ro bảo mật, riêng tư.",
        "Chương V mô tả môi trường cài đặt và kết quả thực nghiệm: detector baseline và đối chứng, tracker, cảnh báo, tốc độ xử lý, cùng phân tích lỗi theo lớp và cỡ vật thể.",
        "Chương VI tổng kết mức độ đạt mục tiêu, các hạn chế về dữ liệu và phương pháp, và đề xuất hướng phát triển."]:
        c.para(t)


# ============================================================ 3.2 (cập nhật W3)
def s32(c):
    K = a.kaggle_w2
    c.heading3("3.2.1 Nguồn dữ liệu và dataset card")
    c.para("Dữ liệu chính là bộ Drone-detection-dataset của Svanström và cộng sự [11], tải thủ công ngày 04/10/2026, giấy phép CC0-1.0 theo repository. Toàn vẹn dữ liệu trên Kaggle được kiểm tra: 1.303/1.303 tệp khớp SHA-256 với bản W1 (checksum_verification.csv). Video được giải mã bằng OpenCV [8].")
    c.table("Dataset card tóm tắt.", ["Thông tin", "Giá trị (nguồn)"],
            [["Số video có nhãn", "650: hồng ngoại 365 (320×256), khả kiến 285 (640×512), 30 fps (video_inventory.csv)"],
             ["Số frame giải mã", "203.328 (stage_status.json, Kaggle run 01)"],
             ["Lớp (số video)", "Airplane 133, Bird 130, Drone 271, Helicopter 116"],
             ["Nhãn", "MATLAB *_LABELS.mat đọc bằng mcos-decoder; mỗi frame tối đa 1 hộp (x, y, w, h); không có track ID"],
             ["Metadata", "Video_dataset_description.xlsx: cảm biến, lớp, số thứ tự, khoảng cách (CLOSE 133 / MEDIUM 292 / DISTANT 225)"],
             ["Phần cứng xử lý", "Kaggle CPU 4 lõi, Python 3.13, OpenCV 4.14 (env.json)"]],
            [3.8, 12.4])
    c.heading3("3.2.2 Thống kê lớp và kích thước vật thể")
    c.image(K / "fig_videos_by_modality_class.png", 12, "Số video theo modality và lớp (Kaggle run 01).")
    c.table("Thống kê kích thước hộp bao (diện tích px²; nhỏ: dưới 32×32 = 1.024 px²; class_stats.csv).",
            ["Modality", "Lớp", "Video", "Frame", "Tỷ lệ hộp nhỏ", "Diện tích trung vị"],
            [["IR", "Airplane", 74, "23.243", "96,4%", "197"], ["IR", "Bird", 79, "24.854", "100%", "178"],
             ["IR", "Drone", 157, "49.043", "98,9%", "219"], ["IR", "Helicopter", 55, "16.955", "86,3%", "381"],
             ["V", "Airplane", 59, "18.585", "45,2%", "1.269"], ["V", "Bird", 51, "16.118", "82,3%", "702"],
             ["V", "Drone", 114, "35.496", "60,5%", "900"], ["V", "Helicopter", 61, "19.034", "22,7%", "4.804"]],
            [2.4, 3.0, 2.2, 2.6, 3.2, 3.2])
    c.image(K / "fig_bbox_area_by_class.png", 14, "Phân bố diện tích hộp bao theo lớp và modality (Kaggle run 01).")
    c.heading3("3.2.3 Chất lượng nhãn")
    c.para("Audit tự động trên toàn bộ 650 video (docs/w2/annotation_audit_w2.md) cho kết quả:")
    c.bullets(["Không có hộp nằm ngoài ảnh, có chiều rộng/cao ≤ 0 hoặc giá trị NaN.",
               "253 video có số mục nhãn khác số frame; 231 video dư đúng một mục. Nhóm này tập trung theo lớp: toàn bộ 114 video Drone khả kiến, 40 video Bird khả kiến và 77 video Bird hồng ngoại. Vì vậy không thể loại bỏ nhóm này mà không mất lớp Drone (ADR-0004).",
               "58 video có bước nhảy tâm hộp lớn; một phần do camera xoay hoặc zoom, nên giả định camera cố định không đúng cho toàn bộ dữ liệu."])
    c.para("Để quyết định cách dùng video dư một mục, nhóm đo độ nhạy: nếu tỷ lệ cặp frame liên tiếp có IoU ≥ 0,5 đạt từ 90% trở lên thì việc lệch một frame hầu như không làm đổi kết quả khớp ở ngưỡng IoU 0,5 của mAP@0.5. Kết quả ở Hình 4: Drone dư một mục đạt 98,6% (95 video) và Bird dư một mục đạt 95,3% (34 video), nên nhóm này được giữ lại với cách căn frame i ↔ nhãn i (offset 0). Nhóm Bird thiếu nhãn chỉ đạt 87,5% và bị loại.")
    c.image(a.fig_w3 / "fig_w3_offset_sensitivity.png", 14, "Độ nhạy với lệch một khung hình theo nhóm video khả kiến (Kaggle run 03a, offset_sensitivity_summary.json).")
    c.para("Nhóm cũng thăm dò chiều lệch bằng bảng xem có nhóm đối chứng, với luật chấm đặt trước (ADR-0004 mục 6). Ở nhóm đối chứng, người xem chỉ chọn đúng 53% (8/15 lựa chọn rõ ràng), thấp hơn ngưỡng 80%, nên phương pháp xem không đủ tin cậy và chiều lệch được ghi là chưa xác định. Kết quả này chỉ là thăm dò, không phải review mù độc lập, vì khóa đáp án nằm cùng gói kết quả với bảng xem và lựa chọn không được chốt trước khi khóa có thể truy cập. Việc kiểm chứng chuyển sang tuần 4 bằng tập chuẩn vẽ lại hộp trên CVAT, với khóa đáp án tách khỏi gói của người xem.")
    c.para("Sau khi phân tích baseline A (mục 5.4), nhóm ghi nhận thêm một chi tiết cấu trúc: 111/114 video Drone khả kiến dư một mục có đúng một mục nhãn rỗng (annotation_audit.csv); ở cả 92 video Drone đã xuất, mục rỗng đó rơi vào frame 0 (yolo_manifest.csv.gz), trong khi review v2 thấy drone ngay ở frame 0. Chi tiết này nghiêng về giả thuyết mục đầu tiên là mục thừa (nhãn đúng của frame i là mục i + 1), nhưng chưa được dùng để đổi quyết định trước khi kiểm chứng ở tuần 4.")
    c.para("Review nhãn lần 2 xem 358 video, mỗi video 8 frame (trải đều và frame có bước nhảy lớn nhất) kèm ảnh phóng to quanh hộp, có hỗ trợ công cụ. Quy tắc: SAI_NHAN khi hộp rỗng trong lúc vật thể nhìn rõ ở chỗ khác hoặc một đoạn dài thiếu nhãn dù vẫn thấy vật thể; KHONG_RO khi không đủ thông tin. Kết quả nêu ở Bảng 10.")
    c.table("Kết quả review nhãn v2 (review_v2_summary.json).", ["Kết luận", "Số video", "Ghi chú"],
            [["OK", 283, "Bao gồm trường hợp lệch một frame ngay tại điểm chuyển cảnh"],
             ["SAI_NHAN", 60, "Tập trung ở lớp Bird; 10 video \"không có nhãn\" thực tế có vật thể bay → không dùng làm ảnh âm tính"],
             ["KHONG_RO", 15, "Vật thể quá mờ hoặc khung hình không đủ thông tin"]],
            [3.0, 2.4, 10.8])
    c.heading3("3.2.4 Chia train/val/test")
    c.para("Đơn vị chia là video (tệp video và nhãn luôn cùng một tập), phân tầng theo modality × lớp, tỷ lệ 70/15/15, seed chọn bằng tìm kiếm tất định (cơ sở 20261011, chọn 20261217) để cân bằng thêm theo số frame, số hộp và khoảng cách. Không chia theo frame. Tăng cường dữ liệu chỉ áp dụng trên train; ngưỡng chọn trên val; test khóa tới đánh giá cuối.")
    c.table("Split v1 (data/splits/video_split_v1.csv).", ["Tập", "Video", "Frame", "Frame có hộp"],
            [["train", 454, "142.187", "136.013"], ["val", 98, "30.545", "28.943"], ["test", 98, "30.596", "29.596"]], [3, 3, 4, 4])
    c.image(a.fig_w2 / "fig_split_v1_by_class.png", 11, "Số video theo lớp trong từng tập của split v1.")
    c.para("Nguy cơ rò rỉ còn lại: 562 cặp video khác tập có khung hình gần giống theo dHash (437 cặp cùng lớp) do dữ liệu thu ở ít vị trí; 123/279 cặp video hồng ngoại – khả kiến cùng lớp và số thứ tự nằm khác tập nhưng chưa có bằng chứng là cùng sự kiện. Vì vậy báo cáo không khẳng định tập test là \"cảnh chưa thấy\".")
    c.heading3("3.2.5 Dữ liệu huấn luyện cho baseline A (tuần 3)")
    c.para("Theo ADR-0004, baseline A chỉ dùng ảnh khả kiến (640×512), bốn lớp Airplane = 0, Bird = 1, Drone = 2, Helicopter = 3, lấy 1 frame trên 10 frame, chỉ xuất train và val. Quá trình chọn video và kết quả xuất được thể hiện ở Hình 6, Hình 7 và Bảng 12.")
    c.image(a.fig_w3 / "fig_w3_data_funnel.png", 15, "Từ video khả kiến tới tập huấn luyện e1-yolo-v1 (Kaggle run 03b, inclusion_decisions.csv).")
    c.image(a.fig_w3 / "fig_w3_yolo_images.png", 15.5, "Số ảnh theo lớp và cỡ hộp bao trong tập e1-yolo-v1 (dataset_stats.csv).")
    c.table("Tập e1-yolo-v1 (dataset_card.json, fingerprint d0ac7e23…).", ["Tập", "Video", "Ảnh", "Ảnh có hộp", "Ảnh không nhãn", "Hộp nhỏ / vừa / lớn"],
            [["train", 178, "5.667", "5.544", 123, "3.041 / 1.894 / 609"], ["val", 38, "1.202", "1.162", 40, "475 / 574 / 113"]],
            [2.0, 2.0, 2.4, 2.6, 2.6, 4.6],
            note="Giới hạn: tập val chỉ còn 3 video Bird (73 ảnh có hộp) nên AP lớp Bird trên val dao động lớn; 77/242 video khả kiến thuộc train và val chưa được xem bằng mắt. Ảnh không nhãn chủ yếu là frame 0 của video dư một mục và đoạn thiếu nhãn, không phải ảnh nền trời trống (mục 5.4).")


# ============================================================ 5.1, 5.3, 5.4 (kết quả W3)
EV4 = "results/W03/kaggle_run_04/train_evidence"
EV5 = "results/W03/kaggle_run_05/analysis"


def s51(c):
    c.para("Mọi bước xử lý dữ liệu thật và huấn luyện chạy trên Kaggle; mỗi lần chạy ghi phiên bản thư viện, phần cứng, commit mã nguồn và SHA-256 của dữ liệu đầu vào. Bảng 13 tóm tắt môi trường của lần huấn luyện baseline A.")
    c.table("Môi trường huấn luyện baseline A (env.json, stage_status.json của Kaggle run 04).", ["Thành phần", "Giá trị"],
            [["Phần cứng", "1 × NVIDIA Tesla T4 15 GB (phiên Kaggle có 2 GPU, Ultralytics dùng GPU 0)"],
             ["Phần mềm", "Python 3.13.15; PyTorch 2.11.0+cu128 [9]; Ultralytics 8.4.183 [12]; OpenCV 4.14 [8]"],
             ["Mã nguồn", "Commit 9597fce (scripts/kaggle/w3_train_baseline.py, configs/baseline_A.yaml)"],
             ["Dữ liệu", "e1-yolo-v1, fingerprint d0ac7e23…, đóng gói một tệp zip, kiểm SHA-256 từng ảnh trước khi huấn luyện"],
             ["Tái lập", "seed 20261013, deterministic = true; weights best.pt/last.pt lưu Google Drive kèm SHA-256"]],
            [3.4, 12.8])


def s53(c):
    c.para("Baseline A là YOLOv8n khởi tạo từ trọng số COCO, ảnh vào 640, tối đa 50 epoch, dừng sớm khi fitness trên val không cải thiện sau 15 epoch (giao thức: docs/w3/training_protocol.md). Toàn bộ số liệu dưới đây đo trên tập validation; tập test vẫn khóa. Vì val vừa dùng để dừng sớm và chọn checkpoint, các số này lạc quan và không được so trực tiếp với chỉ tiêu trên test ở Bảng 1.")
    c.para("Quá trình huấn luyện dừng ở epoch 36, checkpoint tốt nhất ở epoch 21, tổng thời gian khoảng 27 phút (Hình 8). Loss trên train giảm đều trong khi loss hộp trên val không giảm thêm sau khoảng epoch 12–20 và mAP trên val đạt đỉnh ở epoch 21, cho thấy mô hình bắt đầu khớp quá mức với tập train; cơ chế dừng sớm giữ checkpoint tại đỉnh.")
    c.image(a.fig_w3 / "fig_w3_learning_curves.png", 15.5, "Learning curve của baseline A: box loss trên train và val, mAP trên val theo epoch (train/results.csv, Kaggle run 04).")
    c.para("Kết quả theo lớp ở Bảng 14 và Hình 9. AP@0.5 và AP@0.5:0.95 tính trên toàn bộ đường precision–recall với ngưỡng confidence 0,001. Cột P và R là giá trị Ultralytics báo tại một ngưỡng confidence chung cho mọi lớp, chọn để F1 trung bình các lớp lớn nhất; đây không phải ngưỡng vận hành.")
    c.table("Kết quả baseline A trên val (baseline_A_val_metrics.json).", ["Lớp", "P", "R", "AP@0.5", "AP@0.5:0.95"],
            [["Airplane", "0,764", "0,672", "0,677", "0,306"], ["Bird", "0,729", "0,849", "0,865", "0,382"],
             ["Drone", "0,838", "0,871", "0,809", "0,266"], ["Helicopter", "0,921", "0,867", "0,879", "0,414"],
             ["Tất cả", "0,813", "0,815", "0,807", "0,342"]], [3.4, 2.6, 2.6, 3.4, 3.4])
    c.image(a.fig_w3 / "fig_w3_per_class_ap.png", 13.5, "AP@0.5 và AP@0.5:0.95 theo lớp trên val (baseline_A_val_per_class.csv).")
    c.para("Ngưỡng vận hành được chọn theo giao thức: duyệt confidence 0,05–0,90 bước 0,05, chọn ngưỡng có F1 lớp Drone lớn nhất trên val, đồng hạng thì lấy ngưỡng cao hơn. Ngưỡng chọn được là 0,25 (Bảng 15). Ngưỡng có F1 trung bình 4 lớp lớn nhất là 0,35 (F1 = 0,814), chỉ dùng để tham khảo.")
    c.table("Ngưỡng vận hành chọn trên val (operating_points.csv, Kaggle run 05).", ["Confidence", "Drone P", "Drone R", "Drone F1", "4 lớp P", "4 lớp R", "F1 trung bình"],
            [["0,20", "0,794", "0,884", "0,836", "0,786", "0,843", "0,799"],
             ["0,25 (chọn)", "0,832", "0,882", "0,856", "0,814", "0,837", "0,808"],
             ["0,30", "0,834", "0,878", "0,855", "0,826", "0,826", "0,811"],
             ["0,35", "0,837", "0,870", "0,853", "0,835", "0,818", "0,814"]],
            [2.6, 2.0, 2.0, 2.0, 2.0, 2.0, 2.6], size=8.5)


def s54(c):
    c.para("Phân tích lỗi thực hiện tại ngưỡng 0,25 trên val, khớp dự đoán với nhãn khi cùng lớp và IoU ≥ 0,5 (scripts/kaggle/w3_analyze_baseline.py, Kaggle run 05). Bốn nhận xét chính:")
    c.rich([("Bỏ sót chiếm phần lớn lỗi, nhầm lớp ít. ", True), ("Khi bỏ điều kiện cùng lớp, số hộp bị bỏ sót là Airplane 83/289, Drone 58/511, Helicopter 30/285, Bird 2/73; số hộp bị gán sai lớp chỉ từ 1 đến 6 ở mỗi cặp lớp (class_confusion.csv).", False)])
    c.rich([("Recall không giảm đều theo cỡ vật thể. ", True), ("Hộp nhỏ hơn 32×32 px có recall 0,863, không thấp hơn hộp vừa (0,803) (Hình 10). Recall thấp của Airplane cỡ vừa (0,563) tập trung ở ba video: V_AIRPLANE_027 (0/33), V_AIRPLANE_007 (8/30), V_AIRPLANE_005 (16/33) (recall_by_video.csv), nên đây là lỗi theo cảnh hơn là theo kích thước. Helicopter cỡ nhỏ có recall thấp nhất (0,667, 60 hộp).", False)])
    c.image(a.fig_w3 / "fig_w3_recall_by_size.png", 15, "Recall theo lớp và cỡ hộp tại confidence 0,25 trên val; bỏ các ô có dưới 5 hộp (recall_by_size.csv, Kaggle run 05).")
    c.rich([("Định vị hộp còn kém với vật thể nhỏ. ", True), ("Khoảng cách giữa AP@0.5 (0,809) và AP@0.5:0.95 (0,266) lớn nhất ở lớp Drone, nghĩa là mô hình tìm được vật thể nhưng hộp chưa khớp chặt. Một nguyên nhân có thể là nhãn lệch một frame ở video dư một mục: lệch này hầu như không đổi kết quả ở IoU 0,5 (mục 3.2.3) nhưng làm giảm IoU ở các ngưỡng cao.", False)])
    c.rich([("Tập val không có ảnh nền trống thật. ", True), ("Cả 40 ảnh val không có nhãn đều có dự đoán ở ngưỡng 0,25 (negatives_fp.json), nhưng đây không phải dự đoán sai trên nền trời: 20 ảnh là frame 0 của video dư một mục (17 Drone, 3 Bird) — nơi review thấy vật thể nhưng mục nhãn rỗng — và 20 ảnh là đoạn thiếu nhãn f110–f300 của V_BIRD_041. Vì vậy chưa đo được tỷ lệ cảnh báo giả trên nền trời; việc này cần video không có vật thể (tuần 7, hard negatives).", False)])
    c.para("Từ các nhận xét trên, tuần 4 ưu tiên: (1) kiểm chứng chiều lệch nhãn bằng tập chuẩn CVAT và bằng so IoU của dự đoán baseline A với nhãn i và i + 1 trước khi xuất lại dữ liệu; (2) cấu hình đối chứng và tăng cường dữ liệu cho vật thể nhỏ theo ADR-0005; (3) xem lại ba video Airplane có recall thấp để phân biệt lỗi mô hình với lỗi nhãn.")


fill(57, mo_dau); fill(63, s11); fill(66, s12); fill(69, s13); fill(72, s14); fill(75, s15); fill(96, s32); fill(123, s51); fill(129, s53); fill(132, s54)

# Mục lục, danh sách hình, danh sách bảng: trường TOC thật
replace_with_field(40, 'TOC \\o "1-3" \\h \\z \\u', "Nhấn F9 trong Word để cập nhật mục lục.")
replace_with_field(46, 'TOC \\h \\z \\c "Hình"', "Nhấn F9 để cập nhật danh sách hình.")
replace_with_field(51, 'TOC \\h \\z \\c "Bảng"', "Nhấn F9 để cập nhật danh sách bảng.")

# Bảng viết tắt
t = d.tables[2]
extra = [("ADR", "Architecture Decision Record – bản ghi quyết định kỹ thuật"), ("AP", "Average Precision"),
         ("HOTA", "Higher Order Tracking Accuracy"), ("IDF1", "ID F1 score"), ("IoU", "Intersection over Union"),
         ("IR", "Infrared – hồng ngoại"), ("mAP", "mean Average Precision"), ("MVP", "Minimum Viable Product – sản phẩm khả dụng tối thiểu"),
         ("ROI", "Region of Interest – vùng giám sát"), ("UAV", "Unmanned Aerial Vehicle – phương tiện bay không người lái")]
last = t.rows[-1]; last.cells[0].text, last.cells[1].text = extra[0]
for k, v in extra[1:]:
    t._tbl.append(copy.deepcopy(last._tr)); row = t.rows[-1]; row.cells[0].text = k; row.cells[1].text = v

# Tài liệu tham khảo (IEEE)
REFS = [
    "[1] Ultralytics, “Multi-object tracking with Ultralytics YOLO,” Ultralytics Docs. [Online]. Available: https://docs.ultralytics.com/modes/track/ (accessed Oct. 21, 2026).",
    "[2] Y. Zhang, P. Sun, Y. Jiang, D. Yu, F. Weng, Z. Yuan, P. Luo, W. Liu, and X. Wang, “ByteTrack: Multi-object tracking by associating every detection box,” in Proc. Eur. Conf. Comput. Vis. (ECCV), 2022. [Online]. Available: https://arxiv.org/abs/2110.06864",
    "[3] N. Aharon, R. Orfaig, and B.-Z. Bobrovsky, “BoT-SORT: Robust associations multi-pedestrian tracking,” arXiv preprint arXiv:2206.14651, 2022.",
    "[4] Y. Cao et al., “VisDrone-DET2021: The vision meets drone object detection challenge results,” in Proc. IEEE/CVF Int. Conf. Comput. Vis. Workshops (ICCVW), 2021.",
    "[5] G. Chen et al., “VisDrone-MOT2021: The vision meets drone multiple object tracking challenge results,” in Proc. IEEE/CVF Int. Conf. Comput. Vis. Workshops (ICCVW), 2021.",
    "[6] Y. Zhao et al., “DETRs beat YOLOs on real-time object detection,” in Proc. IEEE/CVF Conf. Comput. Vis. Pattern Recognit. (CVPR), 2024. [Online]. Available: https://arxiv.org/abs/2304.08069",
    "[7] D. Du et al., “The unmanned aerial vehicle benchmark: Object detection and tracking,” in Proc. Eur. Conf. Comput. Vis. (ECCV), 2018. [Online]. Available: https://arxiv.org/abs/1804.00518",
    "[8] OpenCV, “OpenCV 4.x documentation.” [Online]. Available: https://docs.opencv.org/4.x/ (accessed Oct. 21, 2026).",
    "[9] PyTorch, “PyTorch documentation.” [Online]. Available: https://pytorch.org/docs/stable/index.html (accessed Oct. 21, 2026).",
    "[10] NVIDIA, “TensorRT documentation.” [Online]. Available: https://docs.nvidia.com/deeplearning/tensorrt/ (accessed Oct. 21, 2026).",
    "[11] F. Svanström, F. Alonso-Fernandez, and C. Englund, “A dataset for multi-sensor drone detection,” Data in Brief, vol. 39, Art. no. 107521, Dec. 2021. Dataset: https://github.com/DroneDetectionThesis/Drone-detection-dataset",
    "[12] G. Jocher, J. Qiu, and A. Chaurasia, “Ultralytics YOLO,” 2023. [Online]. Available: https://github.com/ultralytics/ultralytics (AGPL-3.0).",
]
ref_ps = [p for p in d.paragraphs if p.text[:4].startswith("[") and p.text[1:2].isdigit()]
guide = next(p for p in d.paragraphs if p.text.startswith("[Dùng một chuẩn thống nhất"))
anchor = ref_ps[-1]
for i, p in enumerate(ref_ps):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.add_run(REFS[i])
cur = Cursor(anchor)
for txt in REFS[len(ref_ps):]:
    q = cur.para(txt); q.style = anchor.style
    if anchor._p.pPr is not None:
        q._p.remove(q._p.get_or_add_pPr()); q._p.insert(0, copy.deepcopy(anchor._p.pPr))
guide._p.getparent().remove(guide._p)

# Word tự cập nhật trường khi mở
st = d.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); st.append(uf)
a.out.parent.mkdir(parents=True, exist_ok=True)
d.save(a.out); print("saved", a.out, Cursor.n)
